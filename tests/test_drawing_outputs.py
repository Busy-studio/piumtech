"""API-free rendering and mocked premium tests; no paid image calls."""
import ast
from io import BytesIO
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import patch
import base64

import fitz
import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from smk_drawing_assets import (
    DRAWING_POLICY_VERSION, reset_document_state, make_drawing_slot_reference,
    composite_verified_drawing, reference_manifest, drawing_slot_instructions,
    _box, PREMIUM_DRAWING_SLOT, STYLE_DRAWING_RECT, SLOT_COLOR,
)


def load_app_helpers():
    # Execute definitions, constants and imports, but not the Streamlit UI.
    # Production uses real Streamlit; tests do not need a server or API key.
    path = Path(__file__).resolve().parents[1] / 'app.py'
    tree = ast.parse(path.read_text(encoding='utf-8').split('# Streamlit UI')[0])
    tree.body = [n for n in tree.body if isinstance(n, (ast.Import, ast.ImportFrom, ast.FunctionDef, ast.Assign, ast.AnnAssign))]
    stub = types.ModuleType('streamlit')
    stub.cache_resource = lambda function: function
    stub.warning = lambda *args, **kwargs: None
    stub.secrets = {}
    openai_stub = types.ModuleType('openai'); openai_stub.OpenAI = object
    module = types.ModuleType('app_drawing_tests'); module.__file__ = str(path)
    with patch.dict(sys.modules, {'streamlit': stub, 'openai': openai_stub}):
        exec(compile(tree, str(path), 'exec'), module.__dict__)
    return module


DATA = {
    'language': 'ko', 'university': '테스트대학교', 'department': '테스트학과', 'professor': '홍길동',
    'marketing_title': '대표도면 검증용 기술', 'original_title': '도면 추출 회귀 테스트',
    'subtitle': '검증 데이터', 'applications': [{'name': '산업 제어', 'description': '적용 예시'}],
    'overview': ['원본 도면을 확인합니다.'], 'differentiation': ['다른 문서의 그림을 사용하지 않습니다.'],
    'limitations': ['기존 도면 오선택'], 'technical_advantages': ['원본 기반 도면 선택'],
    'ip': {'title': '검증용 기술', 'application_number': '10-2026-0000001', 'application_date': '2026.10.13'},
    'market_info': {'market_name': '테스트 시장', 'start_year': 2025, 'end_year': 2030,
                    'start_value': 100, 'end_value': 150, 'cagr': 8.4, 'unit': 'USD billion'},
}


class AssetTests(unittest.TestCase):
    def test_reference_sanitization(self):
        source = Image.new('RGB', (1000, 1500), 'black')
        result = make_drawing_slot_reference(source)
        x0, y0, x1, y1 = _box(result, STYLE_DRAWING_RECT)
        region = np.array(result)[y0:y1, x0:x1]
        self.assertTrue(np.all((region == 255).all(2) | (region == SLOT_COLOR).all(2)))
        self.assertEqual(source.getpixel((100, 1100)), (0, 0, 0))

    def test_blank_slot_is_actually_blank(self):
        source = make_drawing_slot_reference(Image.new('RGB', (1000, 1500), 'white'))
        result = composite_verified_drawing(source, None)
        self.assertTrue(np.all(np.array(result) == 255))

    def test_exact_pixels_inserted_and_no_outside_changes(self):
        source = make_drawing_slot_reference(Image.new('RGB', (1000, 1500), (245, 247, 249)))
        drawing = Image.new('RGB', (100, 60), (0, 0, 0))
        result = composite_verified_drawing(source, drawing)
        self.assertEqual(result.getpixel((222, 1113)), (0, 0, 0))
        self.assertEqual(result.getpixel((700, 1000)), source.getpixel((700, 1000)))
        arr = np.array(result)
        self.assertFalse(np.any((arr[:, :, 0] > 220) & (arr[:, :, 1] < 45) & (arr[:, :, 2] > 220)))

    def test_invalid_slot_is_rejected(self):
        for image in (Image.new('RGB', (1000, 1500), 'white'), Image.new('RGB', (1000, 1500), SLOT_COLOR)):
            with self.assertRaises(ValueError): composite_verified_drawing(image, None)

    def test_polluted_slot_is_rejected(self):
        image = make_drawing_slot_reference(Image.new('RGB', (1000, 1500), 'white'))
        ImageDraw.Draw(image).rectangle((60, 1040, 350, 1130), fill='black')
        with self.assertRaises(ValueError): composite_verified_drawing(image, None)

    def test_manifest_matches_optional_assets(self):
        descriptions = ['style', 'palette', 'PIUM/QR']
        text = reference_manifest(descriptions)
        self.assertEqual(text.splitlines(), ['- Image 1: style', '- Image 2: palette', '- Image 3: PIUM/QR'])
        self.assertIn('NO verified drawing', drawing_slot_instructions(False))

    def test_reset_all_outputs_even_on_same_name(self):
        state = dict(data={'old': True}, rep_img=object(), qr_img=object(), hq_png_bytes=b'old',
                     pdf_bytes=b'old', pptx_bytes=b'old', brief_image=object(), source_pdf_sha256='old',
                     edit_version=4, drawing_meta={'old': True}, app_imgs=[object()])
        reset_document_state(state, 'new_bytes_hash')
        for key in ('data', 'rep_img', 'qr_img', 'pdf_bytes', 'pptx_bytes', 'hq_png_bytes', 'brief_image', 'drawing_meta'):
            self.assertIsNone(state[key])
        self.assertEqual(state['app_imgs'], [])
        self.assertEqual(state['source_pdf_sha256'], 'new_bytes_hash')
        self.assertEqual(state['drawing_policy_version'], DRAWING_POLICY_VERSION)
        self.assertEqual(state['edit_version'], 5)


class RenderingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = load_app_helpers()
        # Keep this test deterministic and cheap; market fetch/generation is unrelated.
        cls.app.get_market_visual = lambda *a, **k: (Image.new('RGB', (200, 150), 'white'), 'test')
        cls.app.generate_market_bar_chart = lambda *a, **k: Image.new('RGB', (200, 150), 'white')
        cls.icons = [Image.new('RGB', (100, 100), 'white')]

    def test_basic_pdf_none(self):
        image = self.app.compose_tech_brief(DATA, None, self.icons, '검증용 문의처')
        region = np.asarray(image)[984:1194, 67:370]
        self.assertTrue(np.all(region == 255))
        pdf = self.app.make_pdf_bytes_from_image(image)
        with fitz.open(stream=pdf, filetype='pdf') as doc:
            self.assertEqual(len(doc), 1)
            self.assertGreater(doc[0].get_pixmap().width, 0)

    def test_pptx_none_has_no_drawing_picture(self):
        without = Presentation(BytesIO(self.app.make_pptx_bytes(DATA, None, self.icons, '검증')))
        drawing = Image.new('RGB', (150, 150), 'black')
        with_drawing = Presentation(BytesIO(self.app.make_pptx_bytes(DATA, drawing, self.icons, '검증')))
        count = lambda p: sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in p.slides[0].shapes)
        self.assertEqual(count(with_drawing), count(without) + 1)

    def test_rule_based_premium_none(self):
        image = self.app.make_high_quality_infographic_image(DATA, None, self.icons, '검증')
        self.assertGreater(image.width, 1000)
        self.assertTrue(np.all(np.asarray(image)[2440:2760, 140:910] == 255))

    def test_mocked_premium_manifest_and_exact_blank(self):
        image = make_drawing_slot_reference(Image.new('RGB', (1024, 1536), 'white'))
        bio = BytesIO(); image.save(bio, format='PNG')
        response = types.SimpleNamespace(data=[types.SimpleNamespace(b64_json=base64.b64encode(bio.getvalue()).decode())])
        captured = {}
        def edit(**kwargs): captured.update(kwargs); return response
        fake = types.SimpleNamespace(images=types.SimpleNamespace(edit=edit))
        with patch.object(self.app, 'get_client', return_value=fake):
            result = self.app.generate_reference_based_premium_infographic(DATA, None, self.icons, '검증')
        self.assertTrue(np.all(np.asarray(result) == 255))
        self.assertEqual(len(captured['image']), 3)
        self.assertIn('Image 2: University BRAND COLOR PALETTE', captured['prompt'])
        self.assertIn('Image 3: PIUM/QR', captured['prompt'])
        self.assertNotIn('Image E:', captured['prompt'])
        self.assertTrue(all(handle.closed for handle in captured['image']))
        self.assertTrue(all(not Path(handle.name).exists() for handle in captured['image']))

    def test_mocked_premium_invalid_slot_uses_safe_fallback(self):
        image = Image.new('RGB', (1000, 1500), 'black')
        bio = BytesIO(); image.save(bio, format='PNG')
        response = types.SimpleNamespace(data=[{'b64_json': base64.b64encode(bio.getvalue()).decode()}])
        fake = types.SimpleNamespace(images=types.SimpleNamespace(edit=lambda **kw: response))
        sentinel = Image.new('RGB', (10, 10), 'white')
        with patch.object(self.app, 'get_client', return_value=fake), patch.object(self.app, 'make_high_quality_infographic_image', return_value=sentinel) as fallback:
            result = self.app.generate_reference_based_premium_infographic(DATA, None, self.icons, '검증')
        self.assertIs(result, sentinel)
        self.assertIsNone(fallback.call_args.args[1])


if __name__ == '__main__': unittest.main()
